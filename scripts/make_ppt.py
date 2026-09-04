"""Generates a short slide deck summarizing the project for a mentor
meeting. Run locally (not on Colab): `python scripts/make_ppt.py`
Requires python-pptx (pip install python-pptx).
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DARK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x4A, 0x86, 0xE8)
TEXT = RGBColor(0x2B, 0x2B, 0x2B)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def add_title_bar(slide, kicker, title):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.3))
    tf = box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = kicker.upper()
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = TEXT
    line = slide.shapes.add_shape(1, Inches(0.6), Inches(1.55), Inches(12.1), Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_bullets(slide, items, left=0.7, top=1.9, width=11.9, height=5.0, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, level, bold) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(size if level == 0 else size - 3)
        p.font.bold = bold
        p.font.color.rgb = TEXT if level == 0 else MUTED
        p.space_after = Pt(10)


def bullet(text, level=0, bold=False):
    return (text, level, bold)


# ---------------------------------------------------------------------------
# Slide 1 — Title
# ---------------------------------------------------------------------------
slide = add_slide()
bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = DARK
bg.line.fill.background()

box = slide.shapes.add_textbox(Inches(1), Inches(2.4), Inches(11.3), Inches(2.5))
tf = box.text_frame
tf.word_wrap = True
p0 = tf.paragraphs[0]
p0.text = "Cooperative Highway Merging"
p0.font.size = Pt(40)
p0.font.bold = True
p0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
p1 = tf.add_paragraph()
p1.text = "for Connected Autonomous Vehicles — A CARLA Simulation Study"
p1.font.size = Pt(22)
p1.font.color.rgb = RGBColor(0xCC, 0xCC, 0xDD)
p2 = tf.add_paragraph()
p2.text = "\nProject status review"
p2.font.size = Pt(16)
p2.font.color.rgb = ACCENT

# ---------------------------------------------------------------------------
# Slide 2 — Research question
# ---------------------------------------------------------------------------
slide = add_slide()
add_title_bar(slide, "Motivation", "Research Question & Hypothesis")
add_bullets(slide, [
    bullet("Does explicit cooperation between vehicles improve safety and merge efficiency,", 0, False),
    bullet("compared with each vehicle deciding independently?", 0, False),
    bullet(""),
    bullet("Hypothesis", 0, True),
    bullet("Cooperative yielding / negotiation improves merge success and safety,", 1),
    bullet("while reducing disruptive braking for traffic behind.", 1),
    bullet(""),
    bullet("Scenario", 0, True),
    bullet("One ramp vehicle requests entry; one main-lane vehicle decides:", 1),
    bullet("yield, maintain speed, or accelerate.", 1),
])

# ---------------------------------------------------------------------------
# Slide 3 — Strategies compared
# ---------------------------------------------------------------------------
slide = add_slide()
add_title_bar(slide, "Method", "Five Strategies Compared")
add_bullets(slide, [
    bullet("1. Egoistic — baseline, ignores the ramp vehicle entirely", 0, True),
    bullet("2. Rule-based — yields under a fixed TTC threshold (what most prior work does)", 0, True),
    bullet("3. Negotiation — formal utility model over YIELD / HOLD / ACCELERATE / COUNTER-OFFER", 0, True),
    bullet("weighs delay cost, collision risk, and courtesy — this is the core contribution", 1),
    bullet("4. Negotiation (no uncertainty) — same model, deterministic risk — an ablation", 0, True),
    bullet("5. Learned (stretch goal) — PPO-trained policy, same action vocabulary", 0, True),
    bullet(""),
    bullet("All five share ONE low-level controller (IDM + PID)", 0, True),
    bullet("so differences measure decision quality, not driving skill", 1),
])

# ---------------------------------------------------------------------------
# Slide 4 — What's different / novel
# ---------------------------------------------------------------------------
slide = add_slide()
add_title_bar(slide, "Contribution", "What Makes This Different")
add_bullets(slide, [
    bullet("Risk-aware negotiation under uncertain HDV intent", 0, True),
    bullet("Collision risk = expectation over plausible human-driver behavior,", 1),
    bullet("not a single deterministic point estimate", 1),
    bullet("String-effect metric", 0, True),
    bullet("Speed drop of the SECOND vehicle back — most papers only check the first", 1),
    bullet("Failure-mode taxonomy", 0, True),
    bullet("Classifies why a strategy failed, not just whether it did", 1),
    bullet("Statistical rigor", 0, True),
    bullet("Bootstrap confidence intervals + Kruskal-Wallis / Mann-Whitney significance tests", 1),
])

# ---------------------------------------------------------------------------
# Slide 5 — Honest positioning vs literature
# ---------------------------------------------------------------------------
slide = add_slide()
add_title_bar(slide, "Related Work", "How This Compares to Existing Papers")
add_bullets(slide, [
    bullet("Individual ideas already have precedent — being upfront about this:", 0, True),
    bullet("Bayesian belief over driver intent (Cooperation-Aware Decision Making, IEEE)", 1),
    bullet("Negotiation-based right-of-way assignment (arXiv 1904.06500)", 1),
    bullet("String/disturbance propagation theory (Seiler et al., decades-old control theory)", 1),
    bullet(""),
    bullet("No single paper found combining all of:", 0, True),
    bullet("CARLA + 4-way strategy ladder + shared controller + failure taxonomy +", 1),
    bullet("string-effect metric + uncertainty ablation + proper statistics", 1),
    bullet(""),
    bullet("Contribution = rigor and integration, not a brand-new technique", 0, True),
])

# ---------------------------------------------------------------------------
# Slide 6 — Current status
# ---------------------------------------------------------------------------
slide = add_slide()
add_title_bar(slide, "Status", "Where Things Stand Today")
add_bullets(slide, [
    bullet("Done and verified", 0, True),
    bullet("Negotiation model, metrics, failure classifier, statistics — 40/40 unit tests passing", 1),
    bullet("Full experiment pipeline written: resumable, writes results incrementally", 1),
    bullet(""),
    bullet("In progress", 0, True),
    bullet("Getting the CARLA simulator itself running on Google Colab (no local GPU)", 1),
    bullet("Working through Colab-specific setup issues (Python version, root-user restriction)", 1),
    bullet(""),
    bullet("Not started yet", 0, True),
    bullet("Actual experiment results — zero episodes run so far", 1),
])

# ---------------------------------------------------------------------------
# Slide 7 — Next steps
# ---------------------------------------------------------------------------
slide = add_slide()
add_title_bar(slide, "Plan", "Next Steps")
add_bullets(slide, [
    bullet("1. Finish CARLA-on-Colab setup, confirm one working episode (smoke test)", 0, True),
    bullet("2. Run 1-seed smoke test across all conditions — sanity-check the numbers", 0, True),
    bullet("3. Run the full grid: 5 strategies x 3 densities x 20 seeds = 300 episodes", 0, True),
    bullet("4. Statistical analysis: confidence intervals, significance tests, figures", 0, True),
    bullet("5. Optional: strengthen negotiation model as a formal bargaining game", 0, True),
    bullet("6. Optional: train and evaluate the learned (RL) strategy", 0, True),
    bullet("7. Write up results for submission", 0, True),
])

# ---------------------------------------------------------------------------
# Slide 8 — Limitations (honesty slide)
# ---------------------------------------------------------------------------
slide = add_slide()
add_title_bar(slide, "Limitations", "What This Study Does Not Claim")
add_bullets(slide, [
    bullet("Single-pair interaction (1 ramp + 1 main-lane vehicle), not joint multi-vehicle control", 0),
    bullet("Negotiation utility is a hand-weighted heuristic, not a game-theoretic equilibrium", 0),
    bullet("No results yet — everything above is infrastructure, not findings", 0),
    bullet("Learned strategy is a stretch goal, evaluated only if time permits", 0),
])

prs.save("results/project_overview.pptx")
print("Saved: results/project_overview.pptx")
